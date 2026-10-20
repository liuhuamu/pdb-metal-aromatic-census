#!/usr/bin/env python3
"""Build the supporting-information structure table.

Deliverables
    paper/SI_Table_structures.xlsx
        README      - column definitions + criteria + counts
        Contacts    - all 57,578 metal-aromatic contacts (one row each)
        Structures  - one row per entry (7,953)
    paper/figures/source_data/SI_contacts.csv  (same master table, plain CSV)

Every classification dimension is a filterable column:
geometry class (face-on / edge-on / excluded), bridging class
(bridge-direct / bridge-near / ligand-held / unannotated), metal bucket,
ligand category (DrugBank drug classes / cofactor / additive / other).
"""
from __future__ import annotations

import csv
import json
import sqlite3
from collections import defaultdict
from pathlib import Path

import pandas as pd

ROOT = Path(r"E:\agent\protein\pdb-metal-pi")
PAPER = ROOT / "paper"
BUCKET_FOLDERS = {"A": "A_d10_metpi", "B": "B_alkali_cationpi",
                  "C": "C_transition", "D": "D_noble",
                  "chelated": "chelated"}
AA = ("HIS", "CYS", "ASP", "GLU", "LYS", "ARG", "SER", "THR", "TYR",
      "MET", "ASN", "GLN", "TRP", "PHE", "VAL", "LEU", "ILE", "ALA",
      "GLY", "PRO")

COLUMNS = [
    "pdb_id", "title", "method", "resolution",
    "metal_comp_id", "metal_element", "metal_charge", "metal_bucket",
    "metal_occupancy", "metal_embedded",
    "ligand_comp_id", "ligand_name", "ligand_category",
    "ring_type", "ring_atoms",
    "d_centroid", "d_plane", "offset", "theta_deg", "min_dist",
    "geometry_class", "bridging_class", "coordination",
    "structure_folder",
]


def bridging_class(r) -> str:
    if r["metal_embedded"]:
        return "embedded"
    coord = r["coordination"] or ""
    protein = any(aa in coord for aa in AA)
    if protein and r["min_dist"] <= 2.8:
        return "bridge-direct"
    if protein:
        return "bridge-near"
    if coord.strip():
        return "ligand-held"
    return "unannotated"


def folder_of(r, bridge) -> str:
    if r["geometry_class"] == "excluded":
        return "excluded (not in structure library)"
    return (f"{BUCKET_FOLDERS[r['metal_bucket']]}/"
            f"{r['metal_element']}/{r['geometry_class']}")


def main() -> None:
    # ---- load everything ---------------------------------------------------
    conn = sqlite3.connect(ROOT / "data" / "results.sqlite")
    conn.row_factory = sqlite3.Row
    hits = list(conn.execute("SELECT * FROM hits"))
    conn.close()

    lig_class = {}
    with open(PAPER / "output" / "ligand_classification.csv") as f:
        for row in csv.DictReader(f):
            lig_class[row["comp_id"]] = row["category"]
    lig_names = json.loads((PAPER / "data" / "ligand_names_all.json").read_text())

    # ---- master contact table ----------------------------------------------
    rows = []
    for r in hits:
        bridge = bridging_class(r)
        rows.append({
            "pdb_id": r["pdb_id"],
            "title": r["title"],
            "method": r["method"],
            "resolution": r["resolution"],
            "metal_comp_id": r["metal_comp_id"],
            "metal_element": r["metal_element"],
            "metal_charge": r["metal_charge"],
            "metal_bucket": r["metal_bucket"],
            "metal_occupancy": r["metal_occupancy"],
            "metal_embedded": r["metal_embedded"],
            "ligand_comp_id": r["ligand_comp_id"],
            "ligand_name": lig_names.get(r["ligand_comp_id"], ""),
            "ligand_category": lig_class.get(r["ligand_comp_id"], "other"),
            "ring_type": r["ring_type"],
            "ring_atoms": r["ring_atoms"],
            "d_centroid": r["d_centroid"],
            "d_plane": r["d_plane"],
            "offset": r["offset"],
            "theta_deg": r["theta_deg"],
            "min_dist": r["min_dist"],
            "geometry_class": r["geometry_class"],
            "bridging_class": bridge,
            "coordination": r["coordination"],
            "structure_folder": folder_of(r, bridge),
        })
    df = pd.DataFrame(rows, columns=COLUMNS)

    # ---- per-structure summary ----------------------------------------------
    per = defaultdict(lambda: {"contacts": 0, "metals": set(),
                               "ligands": set(), "classes": set(),
                               "cats": set(), "folders": set(),
                               "title": "", "method": "",
                               "resolution": None})
    for row in rows:
        d = per[row["pdb_id"]]
        d["contacts"] += 1
        d["metals"].add(row["metal_element"])
        d["ligands"].add(row["ligand_comp_id"])
        d["classes"].add(row["geometry_class"])
        d["cats"].add(row["ligand_category"])
        d["folders"].add(row["structure_folder"])
        d["title"] = row["title"]
        d["method"] = row["method"]
        d["resolution"] = row["resolution"]
    df_struct = pd.DataFrame([{
        "pdb_id": pid,
        "title": d["title"],
        "method": d["method"],
        "resolution": d["resolution"],
        "n_contacts": d["contacts"],
        "metals": ";".join(sorted(d["metals"])),
        "ligands": ";".join(sorted(d["ligands"])),
        "geometry_classes": ";".join(sorted(d["classes"])),
        "ligand_categories": ";".join(sorted(d["cats"])),
        "structure_folders": ";".join(sorted(d["folders"])),
    } for pid, d in sorted(per.items())])

    # ---- README sheet ---------------------------------------------------------
    readme = pd.DataFrame({
        "item": [
            "Title",
            "Generated",
            "Contents sheet 'Contacts'",
            "Contents sheet 'Structures'",
            "",
            "COLUMN DEFINITIONS",
            "pdb_id", "title", "method", "resolution",
            "metal_comp_id", "metal_element", "metal_charge",
            "metal_bucket", "metal_occupancy", "metal_embedded",
            "ligand_comp_id", "ligand_name", "ligand_category",
            "ring_type", "ring_atoms",
            "d_centroid", "d_plane", "offset", "theta_deg", "min_dist",
            "geometry_class", "bridging_class", "coordination",
            "structure_folder",
            "",
            "GEOMETRY CLASS CRITERIA (2026-10-16)",
            "face-on", "edge-on", "excluded", "cutoff",
            "",
            "BRIDGING CLASS DEFINITIONS",
            "bridge-direct", "bridge-near", "ligand-held", "embedded",
            "unannotated",
            "",
            "METAL BUCKETS",
            "A", "B", "C", "D",
            "",
            "LIGAND CATEGORIES",
            "cofactor", "additive", "drug-approved", "drug-withdrawn",
            "drug-investigational", "drug-vet_approved",
            "drug-nutraceutical", "drug-experimental", "other",
        ],
        "description": [
            "Metal-aromatic contacts in PDB protein-ligand structures - "
            "supporting table",
            "2026-10-16 (pipeline: pdb-metal-pi)",
            f"all {len(df):,} metal-aromatic ring contacts "
            f"({df.pdb_id.nunique():,} structures)",
            f"one summary row per structure ({len(df_struct):,} rows)",
            "",
            "",
            "PDB entry ID",
            "structure title",
            "experimental method",
            "resolution (A)",
            "metal chemical component ID (e.g. ZN)",
            "metal element",
            "formal charge (from comp dictionary; blank if unknown)",
            "metal bucket A/B/C/D (see below)",
            "metal site occupancy",
            "1 = metal atom embedded in the ligand itself (heme/"
            "chlorophyll class)",
            "ligand chemical component ID",
            "ligand name (RCSB chemcomp)",
            "ligand category (see below)",
            "matching aromatic ring patterns",
            "ring atom names used for centroid/plane",
            "metal...ring centroid distance (A)",
            "perpendicular distance metal...ring plane (A)",
            "lateral offset = sqrt(d_centroid^2 - d_plane^2) (A)",
            "angle metal-centroid vector vs ring normal (deg); the angle "
            "to the ring plane is 90 deg minus this value",
            "closest metal...ring-atom distance (A)",
            "geometry classification: face-on / edge-on / excluded "
            "(criteria below)",
            "metal-mediated binding configuration (definitions below)",
            "metal coordination partners from struct_conn (metalc)",
            "folder in the released structure library containing this "
            "structure's files (excluded contacts have no folder)",
            "",
            "",
            "lateral projection offset < 2.5 A from the ring centre "
            "(the atomic-envelope radius of an aromatic ring)",
            "offset >= 2.5 A, with d_centroid <= 6 A",
            "porphyrin-family ligands (tetrapyrrole/corrin/"
            "phthalocyanine macrocycles) or metal embedded in the ligand",
            "contacts beyond 6 A were not stored",
            "",
            "",
            "metal coordinated by protein residues AND directly bonded to "
            "a ligand ring atom (<=2.8 A)",
            "metal coordinated by protein residues, ligand ring within "
            "6 A",
            "metal held by the ligand itself or water/ions "
            "(e.g. Mg-ATP)",
            "metal is part of the aromatic ligand (heme, chlorophyll)",
            "no struct_conn coordination annotation",
            "",
            "",
            "d10 ions: Zn Cd Hg Cu Ag Au (classic metal-pi)",
            "alkali/alkaline-earth: Li Na K Rb Cs Mg Ca Sr Ba "
            "(classic cation-pi)",
            "other transition metals: Fe Mn Co Ni Cr Mo W V Ti",
            "noble/platinum-group: Pt Pd Ir Rh Ru Os",
            "",
            "",
            "curated list (nucleotides, NAD/FAD, tetrapyrroles, cobalamin, "
            "PQQ, TPP/PLP/folate/CoA/SAM)",
            "curated list (imidazole, benzoic acid, benzamidine, buffers)",
            "DrugBank group: approved",
            "DrugBank group: withdrawn",
            "DrugBank group: investigational",
            "DrugBank group: vet_approved",
            "DrugBank group: nutraceutical",
            "DrugBank group: experimental",
            "not in any list above",
        ],
    })

    # ---- write workbook -----------------------------------------------------
    xlsx = PAPER / "SI_Table_structures.xlsx"
    with pd.ExcelWriter(xlsx, engine="openpyxl") as writer:
        readme.to_excel(writer, sheet_name="README", index=False)
        df.to_excel(writer, sheet_name="Contacts", index=False)
        df_struct.to_excel(writer, sheet_name="Structures", index=False)

        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
        book = writer.book
        widths = {"README": (28, 100), "Contacts": None, "Structures": None}
        for sheet in book.worksheets:
            sheet.freeze_panes = "A2"
            if sheet.title != "README":
                n_cols = sheet.max_column
                sheet.auto_filter.ref = (
                    f"A1:{get_column_letter(n_cols)}{sheet.max_row}")
            if sheet.title == "README":
                for col, w in zip("AB", widths["README"]):
                    sheet.column_dimensions[col].width = w
            else:
                for i in range(1, sheet.max_column + 1):
                    letter = get_column_letter(i)
                    name = str(sheet.cell(1, i).value)
                    if name in ("title", "coordination", "ring_atoms"):
                        sheet.column_dimensions[letter].width = 40
                    elif name in ("ligand_name", "structure_folder",
                                  "ligand_categories", "geometry_classes",
                                  "structure_folders", "ligands"):
                        sheet.column_dimensions[letter].width = 28
                    else:
                        sheet.column_dimensions[letter].width = 13
            for cell in sheet[1]:
                cell.font = Font(bold=True)

    # ---- CSV copy ----------------------------------------------------------
    df.to_csv(PAPER / "figures" / "source_data" / "SI_contacts.csv",
              index=False)

    # ---- console summary ----------------------------------------------------
    print(f"contacts table : {len(df):,} rows x {len(df.columns)} columns")
    print(f"structures table: {len(df_struct):,} rows")
    print("quick lookup counts (Contacts sheet, filter column "
          "geometry_class):")
    for cls, n in df.geometry_class.value_counts().items():
        print(f"  {cls:<9s} {n:>7,}")
    print("bridging_class:")
    for cls, n in df.bridging_class.value_counts().items():
        print(f"  {cls:<14s} {n:>7,}")
    print(f"\nsaved -> {xlsx}")
    print(f"csv   -> paper/figures/source_data/SI_contacts.csv")


if __name__ == "__main__":
    main()
